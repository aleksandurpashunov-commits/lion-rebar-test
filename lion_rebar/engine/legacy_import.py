from __future__ import annotations

import json
import re
from collections import defaultdict
from copy import deepcopy
from pathlib import Path
from typing import Any

from lion_rebar.config.constants import DEFAULT_CATEGORIES, DEFAULT_POSITION_STATUSES
from lion_rebar.engine.models import (
    CategoryDefinition,
    EnginePackage,
    RuleDefinition,
    SegmentDefinition,
    ShapeDefinition,
    SourceAsset,
    utc_now_iso,
)
from lion_rebar.utils.files import read_json, sha256_file


class LegacyImportError(ValueError):
    pass


def extract_json_object_after_marker(text: str, marker: str) -> dict[str, Any]:
    marker_index = text.find(marker)
    if marker_index < 0:
        raise LegacyImportError(f"Не е намерен маркерът {marker!r}.")
    start = text.find("{", marker_index + len(marker))
    if start < 0:
        raise LegacyImportError(f"Не е намерен JSON обект след {marker!r}.")

    depth = 0
    in_string = False
    escaped = False
    for index in range(start, len(text)):
        char = text[index]
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
            continue
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                raw = text[start : index + 1]
                try:
                    value = json.loads(raw)
                except json.JSONDecodeError as exc:
                    raise LegacyImportError(
                        f"Вграденият DEFAULT_LIBRARY не е валиден JSON: {exc}"
                    ) from exc
                if not isinstance(value, dict):
                    raise LegacyImportError("DEFAULT_LIBRARY не е обект.")
                return value
    raise LegacyImportError("Незавършен JSON обект в HTML engine-а.")


def import_legacy_html(path: Path) -> EnginePackage:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    library = extract_json_object_after_marker(text, "const DEFAULT_LIBRARY=")
    version = _find_constant(text, "EDITOR_VERSION") or _version_from_filename(path.name) or "0.1.0"
    reserved_codes = _extract_reserved_codes(text)
    engine = engine_from_legacy_library(
        library,
        version=version,
        source_path=path,
        reserved_codes=reserved_codes,
    )
    engine.source_assets.append(
        SourceAsset(
            role="MASTER_ENGINE_HTML",
            filename=path.name,
            sha256=sha256_file(path),
            notes="Legacy HTML engine imported as structured draft. Original file remains unchanged.",
        )
    )
    engine.extensions["legacy_html"] = {
        "editor_version": version,
        "base_engine_file": _find_constant(text, "BASE_ENGINE_FILE"),
        "base_engine_sha256": _find_constant(text, "BASE_ENGINE_SHA256"),
        "reserved_codes": sorted(reserved_codes),
    }
    return engine


def import_legacy_xlsx(path: Path, version: str | None = None) -> EnginePackage:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise LegacyImportError("openpyxl не е инсталиран.") from exc

    workbook = load_workbook(path, data_only=False, read_only=True)
    required = ["ФОРМИ", "СЕГМЕНТИ", "ИЗЧЕРТАВАНЕ", "ТЕСТОВЕ", "НАСТРОЙКИ", "РАЗМЕРИ"]
    missing = [name for name in required if name not in workbook.sheetnames]
    if missing:
        raise LegacyImportError(f"Липсващи листове в библиотеката: {', '.join(missing)}")

    library: dict[str, list[dict[str, Any]]] = {}
    for sheet_name in required:
        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
        header_index = next(
            (index for index, row in enumerate(rows) if any(value not in (None, "") for value in row)),
            None,
        )
        if header_index is None:
            library[sheet_name] = []
            continue
        headers = [str(value).strip() if value is not None else "" for value in rows[header_index]]
        records: list[dict[str, Any]] = []
        for row in rows[header_index + 1 :]:
            record = {
                header: value
                for header, value in zip(headers, row)
                if header and value not in (None, "")
            }
            if record:
                records.append(record)
        library[sheet_name] = records

    detected_version = version or _version_from_filename(path.name) or "0.1.0"
    reserved_codes = {
        str(row.get("КОД_ФОРМА", "")).upper()
        for row in library.get("ФОРМИ", [])
        if str(row.get("СТАТУС", "")).upper() in {"ЗАКЛЮЧЕН", "ОДОБРЕН"}
    }
    engine = engine_from_legacy_library(
        library,
        version=detected_version,
        source_path=path,
        reserved_codes=reserved_codes,
    )
    engine.source_assets.append(
        SourceAsset(
            role="MASTER_LIBRARY_XLSX",
            filename=path.name,
            sha256=sha256_file(path),
            notes="Structured MASTER library imported from XLSX.",
        )
    )
    return engine


def engine_from_legacy_library(
    library: dict[str, list[dict[str, Any]]],
    *,
    version: str,
    source_path: Path,
    reserved_codes: set[str] | None = None,
) -> EnginePackage:
    reserved_codes = {code.upper() for code in (reserved_codes or set())}
    segment_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    drawing_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    dimension_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    test_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for row in library.get("СЕГМЕНТИ", []):
        segment_rows[str(row.get("КОД_ФОРМА", "")).upper()].append(row)
    for row in library.get("ИЗЧЕРТАВАНЕ", []):
        drawing_rows[str(row.get("КОД_ФОРМА", "")).upper()].append(row)
    for row in library.get("РАЗМЕРИ", []):
        dimension_rows[str(row.get("КОД_ФОРМА", "")).upper()].append(row)
    for row in library.get("ТЕСТОВЕ", []):
        test_rows[str(row.get("КОД_ФОРМА", "")).upper()].append(row)

    shapes: list[ShapeDefinition] = []
    for form_row in library.get("ФОРМИ", []):
        code = str(form_row.get("КОД_ФОРМА", "")).strip().upper()
        if not code:
            continue
        drawings_by_index = {
            _to_int(row.get("№_СЕГМЕНТ"), 0): row for row in drawing_rows.get(code, [])
        }
        dimensions_by_index = {
            _to_int(row.get("№_СЕГМЕНТ"), 0): row for row in dimension_rows.get(code, [])
        }
        segments: list[SegmentDefinition] = []
        for row in sorted(segment_rows.get(code, []), key=lambda item: _to_int(item.get("№_СЕГМЕНТ"), 0)):
            index = _to_int(row.get("№_СЕГМЕНТ"), len(segments) + 1)
            drawing = drawings_by_index.get(index, {})
            dimension = dimensions_by_index.get(index, {})
            segments.append(
                SegmentDefinition(
                    index=index,
                    length_code=str(row.get("КОД_ДЪЛЖИНА", f"L{index}")),
                    segment_type=_map_segment_type(row.get("ТИП_СЕГМЕНТ")),
                    fixed_length_cm=_to_float_or_none(row.get("ФИКСИРАН_РАЗМЕР_СМ")),
                    bend_after_deg=_to_float_or_none(row.get("ЪГЪЛ_СЛЕД_СЕГМЕНТА_ГРАДУСИ")),
                    bend_direction=_map_bend_direction(row.get("ПОСОКА_ОГЪВКА")),
                    geometry_method=str(row.get("МЕТОД_ГЕОМЕТРИЯ", "DIRECT_LENGTH")),
                    projection_x_code=_empty_to_none(row.get("КОД_ПРОЕКЦИЯ_X")),
                    projection_y_code=_empty_to_none(row.get("КОД_ПРОЕКЦИЯ_Y")),
                    radius_cm=_to_float_or_none(row.get("РАДИУС_ОГЪВКА_СМ")),
                    required=_yes_no(row.get("ЗАДЪЛЖИТЕЛЕН", "ДА")),
                    drawing_direction=_empty_to_none(drawing.get("ПОСОКА_ИЗЧЕРТАВАНЕ")),
                    dimension_position=str(dimension.get("ПОЗИЦИЯ_РАЗМЕР", "AUTO")),
                    dimension_offset=_dimension_offset(dimension),
                    notes=str(row.get("БЕЛЕЖКИ", "")),
                    raw={
                        "segment": deepcopy(row),
                        "drawing": deepcopy(drawing),
                        "dimension": deepcopy(dimension),
                    },
                )
            )

        status = str(form_row.get("СТАТУС", "")).upper()
        locked = code in reserved_codes or status in {"ЗАКЛЮЧЕН", "LOCKED"}
        state = "LOCKED" if locked else ("APPROVED" if status in {"ОДОБРЕН", "ПОДГОТВЕН"} else "DRAFT")
        shapes.append(
            ShapeDefinition(
                code=code,
                name=str(form_row.get("ИМЕ_ФОРМА", code)),
                category=str(form_row.get("КАТЕГОРИЯ", "")),
                state=state,
                revision=max(1, _to_int(form_row.get("ВЕРСИЯ"), 1)),
                locked=locked,
                active=True,
                initial_direction=_infer_initial_direction(drawings_by_index.get(1, {})),
                closed=_yes_no(form_row.get("ЗАТВОРЕНА", "НЕ")),
                symmetric=_yes_no(form_row.get("СИМЕТРИЧНА", "НЕ")),
                mirror_allowed=_yes_no(form_row.get("ОГЛЕДАЛНА_ДОПУСТИМА", "НЕ")),
                template_orientation=str(form_row.get("ОРИЕНТАЦИЯ_ПО_ШАБЛОН", "")),
                notes=str(form_row.get("БЕЛЕЖКИ", "")),
                segments=segments,
                rendering={
                    "legacy_drawing_rows": deepcopy(drawing_rows.get(code, [])),
                    "fit_mode": _setting_value(library, "FIT_MODE"),
                    "drawing_mode": _setting_value(library, "DRAWING_MODE"),
                },
                recognition={
                    "segment_count": _to_int(form_row.get("БРОЙ_СЕГМЕНТИ"), len(segments)),
                    "bend_count": _to_int(form_row.get("БРОЙ_ОГЪВКИ"), max(0, len(segments) - 1)),
                },
                dimensions=deepcopy(dimension_rows.get(code, [])),
                legacy_raw={"form": deepcopy(form_row)},
            )
        )

    settings = {
        "max_segments": 12,
        "published_snapshots_are_immutable": True,
        "project_pins_engine_version": True,
        "require_all_tests_to_pass_for_production": True,
        "legacy_settings": {
            str(row.get("КЛЮЧ", "")): row.get("СТОЙНОСТ")
            for row in library.get("НАСТРОЙКИ", [])
            if row.get("КЛЮЧ")
        },
    }
    categories = [
        CategoryDefinition(
            code=item["code"],
            sheet_name=item["sheet_name"],
            aliases=item["aliases"],
            sort_order=index * 10,
        )
        for index, item in enumerate(DEFAULT_CATEGORIES, start=1)
    ]
    rules = [
        RuleDefinition(
            code="R-BEAM-90-F003-F004",
            name="Сложно изчертаните гредови винкели се тълкуват като 90°",
            scope="SHAPE_CLASSIFICATION",
            priority=10,
            enabled=True,
            when={
                "all": [
                    {"field": "element_type", "operator": "equals", "value": "BEAM"},
                    {"field": "engineering_bend", "operator": "equals", "value": 90},
                ]
            },
            then=[
                {"action": "map_one_end_to", "shape": "F004"},
                {"action": "map_two_ends_to", "shape": "F003"},
                {"action": "forbid_auto_new_shape", "value": True},
            ],
            notes="Заключено проектно правило от MASTER prompt. Изпълнението му се добавя в етап 4.",
        )
    ]
    engine = EnginePackage(
        engine_id="lion-rebar-master",
        title=f"Imported LION REBAR MASTER from {source_path.name}",
        version=version,
        state="DRAFT",
        parent_version=None,
        production_ready=bool(shapes),
        shapes=shapes,
        rules=rules,
        categories=categories,
        position_statuses=list(DEFAULT_POSITION_STATUSES),
        settings=settings,
        tests=deepcopy(library.get("ТЕСТОВЕ", [])),
        change_log=[
            {
                "at": utc_now_iso(),
                "action": "IMPORT_LEGACY_MASTER",
                "source": source_path.name,
                "forms": len(shapes),
            }
        ],
    )
    return engine


def import_shape_draft_json(path: Path) -> list[ShapeDefinition]:
    data = read_json(path)
    if isinstance(data, list):
        raw_forms = data
    elif isinstance(data, dict):
        raw_forms = data.get("forms") or data.get("shapes") or [data]
    else:
        raise LegacyImportError("Неподдържан JSON формат за форми.")

    shapes: list[ShapeDefinition] = []
    for raw in raw_forms:
        if not isinstance(raw, dict):
            continue
        code = str(
            raw.get("proposedCode")
            or raw.get("proposed_code")
            or raw.get("code")
            or ""
        ).upper()
        if not code:
            continue
        raw_segments = raw.get("segments", [])
        segments: list[SegmentDefinition] = []
        for index, segment in enumerate(raw_segments, start=1):
            if not isinstance(segment, dict):
                continue
            bend = segment.get("angle", segment.get("bend_after_deg"))
            turn = segment.get("turn", segment.get("bend_direction"))
            segments.append(
                SegmentDefinition(
                    index=index,
                    length_code=f"L{index}",
                    fixed_length_cm=_to_float_or_none(
                        segment.get("length", segment.get("length_cm"))
                    ),
                    bend_after_deg=_to_float_or_none(bend),
                    bend_direction=_map_bend_direction(turn),
                    drawing_direction=None,
                    dimension_position=str(
                        segment.get("dimensionPosition", segment.get("dimension_position", "AUTO"))
                    ),
                    dimension_offset=float(
                        segment.get("dimensionOffset", segment.get("dimension_offset", 6)) or 6
                    ),
                    raw=deepcopy(segment),
                )
            )
        shapes.append(
            ShapeDefinition(
                code=code,
                name=str(raw.get("name", raw.get("ИМЕ_ФОРМА", "НОВА ФОРМА"))),
                category=str(raw.get("category", "")),
                state="DRAFT",
                revision=_to_int(raw.get("revision"), 1),
                locked=False,
                closed=_yes_no(raw.get("closed", raw.get("ЗАТВОРЕНА", False))),
                symmetric=_yes_no(
                    raw.get("symmetric", raw.get("СИМЕТРИЧНА", False))
                ),
                initial_direction=str(
                    raw.get("initialDirection", raw.get("initial_direction", "RIGHT"))
                ),
                mirror_allowed=_yes_no(
                    raw.get(
                        "mirrorAllowed",
                        raw.get("mirror_allowed", raw.get("ОГЛЕДАЛНА_ДОПУСТИМА", False)),
                    )
                ),
                notes=str(raw.get("notes", "")),
                segments=segments,
                rendering={"separate_overlap": bool(raw.get("separateOverlap", True))},
                recognition=deepcopy(raw.get("recognition", {})),
                dimensions=deepcopy(raw.get("dimensions", [])),
                tests=deepcopy(raw.get("tests", [])),
                legacy_raw=deepcopy(raw),
            )
        )
    if not shapes:
        raise LegacyImportError("JSON файлът не съдържа разпознаваема форма.")
    return shapes


def _find_constant(text: str, name: str) -> str | None:
    match = re.search(rf"const\s+{re.escape(name)}\s*=\s*['\"]([^'\"]+)['\"]", text)
    return match.group(1) if match else None


def _version_from_filename(filename: str) -> str | None:
    match = re.search(r"v(\d+)\.(\d+)\.(\d+)", filename, re.IGNORECASE)
    return ".".join(match.groups()) if match else None


def _extract_reserved_codes(text: str) -> set[str]:
    match = re.search(r"RESERVED_CODES\s*=\s*new\s+Set\s*\(\s*\[([^\]]*)\]", text)
    if not match:
        return set()
    return {code.upper() for code in re.findall(r"['\"](F\d+)['\"]", match.group(1), re.IGNORECASE)}


def _yes_no(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().upper() in {"ДА", "YES", "TRUE", "1", "Y"}


def _to_int(value: Any, default: int) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _to_float_or_none(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _empty_to_none(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


def _map_segment_type(value: Any) -> str:
    text = str(value or "").upper()
    if "ДЪГА" in text or "ARC" in text or "КРИВ" in text:
        return "ARC"
    return "LINE"


def _map_bend_direction(value: Any) -> str | None:
    text = str(value or "").strip().upper()
    mapping = {
        "НАЛЯВО": "LEFT",
        "ЛЯВО": "LEFT",
        "LEFT": "LEFT",
        "НАДЯСНО": "RIGHT",
        "ДЯСНО": "RIGHT",
        "RIGHT": "RIGHT",
        "НЯМА": "NONE",
        "NONE": "NONE",
        "": None,
    }
    return mapping.get(text, text or None)


def _dimension_offset(row: dict[str, Any]) -> float:
    x = _to_float_or_none(row.get("ОТМЕСТВАНЕ_X")) or 0.0
    y = _to_float_or_none(row.get("ОТМЕСТВАНЕ_Y")) or 0.0
    value = max(abs(x), abs(y))
    return value if value > 0 else 6.0


def _infer_initial_direction(first_drawing: dict[str, Any]) -> str:
    direction = str(first_drawing.get("ПОСОКА_ИЗЧЕРТАВАНЕ", "НАДЯСНО")).upper()
    mapping = {
        "НАДЯСНО": "RIGHT",
        "НАЛЯВО": "LEFT",
        "НАГОРЕ": "UP",
        "НАДОЛУ": "DOWN",
    }
    return mapping.get(direction, "RIGHT")


def _setting_value(library: dict[str, list[dict[str, Any]]], key: str) -> Any:
    for row in library.get("НАСТРОЙКИ", []):
        if str(row.get("КЛЮЧ", "")).upper() == key.upper():
            return row.get("СТОЙНОСТ")
    return None
