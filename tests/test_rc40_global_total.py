from __future__ import annotations

import base64
from dataclasses import replace
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from lion_rebar.engine.models import (
    EnginePackage,
    ParameterRelation,
    SegmentDefinition,
    ShapeDefinition,
)
from lion_rebar.export.excel_export import (
    ExportGroup,
    ExportFrameConfig,
    ExportPosition,
    _master_topology_signature,
    _production_identity_token,
    _sumif_identity_formula,
    _validate_f002_turn_consistency,
    _validate_global_aggregation,
    _validate_written_global_formulas,
    generate_excel_export,
)
from lion_rebar.export.errors import ExcelExportError
from lion_rebar.export.frame_detection import DrawingFrame
from lion_rebar.project.models import ProjectModel


_PNG_1X1 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/"
    "x8AAusB9Y9Z7r0AAAAASUVORK5CYII="
)


def _engine() -> EnginePackage:
    return EnginePackage(
        engine_id="rc40-test-engine",
        title="RC40 TEST",
        version="1.0",
        content_sha256="rc40-test-sha",
        production_ready=True,
        shapes=[
            ShapeDefinition(
                code="F001",
                name="ПРАВ ПРЪТ",
                state="LOCKED",
                locked=True,
                revision=2,
                segments=[SegmentDefinition(index=1, length_code="L1")],
            )
        ],
    )


def _frame(index: int, category: str, sheet: str) -> DrawingFrame:
    return DrawingFrame(
        index,
        f"FRAME-{index}",
        index * 1000.0,
        0.0,
        sheet,
        "",
        category,
        sheet,
        "HIGH",
    )


def _position(index: int, count: int) -> ExportPosition:
    beam = "Гр. A.1 25/50" if index == 4 else ""
    return ExportPosition(
        annotation_handle=f"A{index}",
        annotation_text=f"{count}N12x200",
        count=count,
        diameter_mm=12,
        production_length_cm=200,
        shape_code="F001",
        shape_name="ПРАВ ПРЪТ",
        shape_revision=2,
        geometry_handles=(f"G{index}",),
        frame_index=index,
        segment_lengths_cm=(200,),
        turn_angles_deg=(),
        beam_name=beam,
        beam_label_handle="BEAM-A1" if beam else "",
    )


def test_total_merges_one_exact_identity_across_every_source_sheet(
    tmp_path: Path, monkeypatch,
) -> None:
    import lion_rebar.export.excel_export as exporter
    import openpyxl.drawing.image as openpyxl_image
    import openpyxl.reader.excel as openpyxl_reader

    monkeypatch.setattr(openpyxl_image, "PILImage", False)
    monkeypatch.setattr(openpyxl_reader, "find_images", lambda *_args, **_kwargs: ([], []))

    def fake_render(path: Path, **_kwargs) -> Path:
        path.write_bytes(base64.b64decode(_PNG_1X1))
        return path

    monkeypatch.setattr(exporter, "render_shape_png", fake_render)

    categories = (
        ("BOTTOM", "долна"),
        ("TOP", "горна"),
        ("DETAILS", "детайли"),
        ("BEAMS", "греди"),
        ("VERTICALS", "вертикали"),
    )
    frames = [
        _frame(index, category, sheet)
        for index, (category, sheet) in enumerate(categories, start=1)
    ]
    configs = [
        ExportFrameConfig(frame.index, category, sheet, frame.combined_title)
        for frame, (category, sheet) in zip(frames, categories)
    ]
    counts = (2, 3, 5, 7, 11)
    positions = [
        _position(index, count)
        for index, count in enumerate(counts, start=1)
    ]
    output = tmp_path / "rc40_all_sheets.xlsx"

    result = generate_excel_export(
        ProjectModel(name="RC40", root=str(tmp_path)),
        _engine(),
        frames,
        positions,
        configs,
        output,
    )

    assert result.total_group_count == 1
    assert set(result.position_numbers_by_annotation.values()) == {1}

    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == [
        "долна", "горна", "детайли", "греди", "вертикали", "Общо"
    ]
    total_formula = workbook["Общо"]["B7"].value
    assert isinstance(total_formula, str) and total_formula.startswith("=SUMIF(")
    for sheet in ("долна", "горна", "детайли", "греди", "вертикали"):
        assert f"'{sheet}'!$T$" in total_formula
    assert workbook["Общо"]["A7"].value == 1
    assert workbook["Общо"]["B7"].value.count("SUMIF(") == 5
    assert workbook["Общо"]["J7"].value.count("SUMIF(") == 5
    assert workbook["Общо"]["K7"].value.count("SUMIF(") == 5
    assert workbook["Общо"]["T7"].value.startswith("pid-sha256:")
    for sheet in ("долна", "горна", "детайли", "греди", "вертикали"):
        assert workbook[sheet]["T7" if sheet != "греди" else "T8"].value == (
            workbook["Общо"]["T7"].value
        )

    trace = json.loads(result.traceability_path.read_text(encoding="utf-8"))
    audit = trace["global_aggregation"]
    expected_source_sheets = sorted(
        ["долна", "горна", "детайли", "греди", "вертикали"]
    )
    assert audit["source_sheets"] == expected_source_sheets
    assert audit["source_sheet_count"] == 5
    assert audit["input_annotation_count"] == 5
    assert audit["source_group_count"] == 5
    assert audit["total_group_count"] == 1
    assert audit["input_count_total"] == sum(counts)
    assert audit["input_count_total"] == audit["source_count_total"]
    assert audit["input_count_total"] == audit["total_count_total"]
    assert audit["input_mass_kg"] == audit["source_mass_kg"]
    assert audit["input_mass_kg"] == audit["total_mass_kg"]
    assert audit["one_total_row_per_identity"] is True
    assert audit["duplicate_identity_key_count"] == 0
    assert audit["formula_strategy"] == (
        "SUMIF_BY_HIDDEN_PRODUCTION_IDENTITY_V1"
    )
    assert audit["formula_cells_validated"] == 3
    assert audit["excel_255_argument_limit_avoided"] is True
    total_row = next(row for row in trace["rows"] if row.get("kind") == "TOTAL")
    assert total_row["source_group_count"] == 5
    assert total_row["source_sheets"] == expected_source_sheets
    assert total_row["master_topology_signature"].startswith("sha256:")


@pytest.mark.parametrize(
    "change",
    [
        {"shape_code": "F005"},
        {"diameter_mm": 12.01},
        {"production_length_cm": 200.1},
        {"segment_lengths_cm": (199.9, 0.1)},
        {"turn_angles_deg": (90.0,)},
        {"shape_revision": 3},
        {"master_topology_signature": "sha256:other"},
    ],
)
def test_global_identity_separates_every_production_dimension(change: dict) -> None:
    base = replace(
        _position(1, 1),
        master_topology_signature="sha256:base",
    )
    assert replace(base, **change).group_key != base.group_key


@pytest.mark.parametrize(
    "outside_change",
    [
        {"diameter_mm": 12.00006},
        {"production_length_cm": 200.06},
        {"segment_lengths_cm": (200.06,)},
        {"turn_angles_deg": (90.06,)},
    ],
)
def test_global_identity_uses_documented_production_precision(
    outside_change: dict,
) -> None:
    base = replace(
        _position(1, 1),
        production_length_cm=200.04,
        segment_lengths_cm=(200.04,),
        turn_angles_deg=(89.96,),
        diameter_mm=12.00001,
        master_topology_signature="sha256:base",
    )
    same_precision = replace(
        base,
        production_length_cm=200.049,
        segment_lengths_cm=(200.049,),
        turn_angles_deg=(89.99,),
        diameter_mm=12.00004,
    )
    assert same_precision.group_key == base.group_key
    assert replace(base, **outside_change).group_key != base.group_key


def test_master_topology_digest_is_sensitive_to_manufacturing_semantics() -> None:
    segment = SegmentDefinition(
        index=1,
        length_code="L1",
        bend_after_deg=90,
        bend_direction="LEFT",
        radius_cm=2,
    )
    relation = ParameterRelation(
        kind="RATIO", left_segment=1, right_segment=2, ratio=2
    )
    shape = ShapeDefinition(
        code="X",
        name="X",
        initial_direction="RIGHT",
        closed=False,
        segments=[segment],
        parameter_relations=[relation],
    )
    base = _master_topology_signature(shape)
    variants = [
        replace(shape, closed=True),
        replace(shape, initial_direction="LEFT"),
        replace(shape, segments=[replace(segment, bend_after_deg=45)]),
        replace(shape, segments=[replace(segment, bend_direction="RIGHT")]),
        replace(shape, segments=[replace(segment, radius_cm=3)]),
        replace(shape, parameter_relations=[replace(relation, ratio=3)]),
    ]
    assert all(_master_topology_signature(item) != base for item in variants)
    assert _master_topology_signature(
        replace(shape, rendering={"color": "red"})
    ) == base


def test_f002_inconsistent_second_turn_fails_closed() -> None:
    valid = replace(
        _position(1, 1),
        shape_code="F002",
        turn_angles_deg=(5.71, -5.72),
    )
    _validate_f002_turn_consistency(valid)
    with pytest.raises(ExcelExportError, match="inconsistent reduction turns"):
        _validate_f002_turn_consistency(
            replace(valid, turn_angles_deg=(5.7, -25.0))
        )


def test_f002_identity_is_invariant_to_reverse_dxf_traversal() -> None:
    forward = replace(
        _position(1, 1),
        shape_code="F002",
        segment_lengths_cm=(50.0, 36.5, 25.0),
        turn_angles_deg=(5.79, -5.91),
        master_topology_signature="sha256:f002",
    )
    reverse = replace(
        forward,
        segment_lengths_cm=(25.0, 36.5, 50.0),
        turn_angles_deg=(-5.91, 5.79),
    )
    _validate_f002_turn_consistency(forward)
    _validate_f002_turn_consistency(reverse)
    assert forward.group_key == reverse.group_key


def test_f002_same_direction_or_non_finite_turns_fail_closed() -> None:
    position = replace(_position(1, 1), shape_code="F002")
    with pytest.raises(ExcelExportError, match="opposite signs"):
        _validate_f002_turn_consistency(
            replace(position, turn_angles_deg=(5.71, 5.72))
        )
    for value in (float("nan"), float("inf"), float("-inf")):
        with pytest.raises(ExcelExportError, match="non-finite"):
            _validate_f002_turn_consistency(
                replace(position, turn_angles_deg=(value, -5.7))
            )


def _written_formula_fixture():
    workbook = Workbook()
    source_sheet = workbook.active
    source_sheet.title = "долна"
    total_sheet = workbook.create_sheet("Общо")
    position = replace(
        _position(1, 2), master_topology_signature="sha256:base"
    )
    source_group = ExportGroup(key=position.group_key, positions=[position])
    total_group = ExportGroup(
        key=position.group_key,
        positions=[position],
        row_number=7,
    )
    source_refs = [("долна", 7, source_group)]
    sources = {position.group_key: source_refs}
    token = _production_identity_token(position.group_key)
    source_sheet["B7"] = "=2"
    source_sheet["D7"] = 12
    source_sheet["F7"] = 200
    source_sheet["I7"] = "=(D7/2)*(D7/2)*3.14*0.007859872"
    source_sheet["J7"] = "=B7*F7/100"
    source_sheet["K7"] = "=J7*I7"
    source_sheet["T7"] = token
    total_sheet["D7"] = 12
    total_sheet["F7"] = 200
    total_sheet["I7"] = "=(D7/2)*(D7/2)*3.14*0.007859872"
    total_sheet["T7"] = token
    for column in ("B", "J", "K"):
        total_sheet[f"{column}7"] = _sumif_identity_formula(
            source_refs,
            total_row=7,
            value_column=column,
        )
    return workbook, total_sheet, sources, [total_group]


def test_written_formula_validator_checks_source_tokens_and_values() -> None:
    workbook, total_sheet, sources, groups = _written_formula_fixture()
    audit = _validate_written_global_formulas(
        workbook, total_sheet, sources, groups
    )
    assert audit["source_identity_cells_validated"] == 1
    assert audit["source_value_cells_validated"] == 6
    assert audit["source_coordinates_unique"] is True


@pytest.mark.parametrize(
    "cell", ["B7", "D7", "F7", "I7", "J7", "K7", "T7"]
)
def test_written_formula_validator_rejects_total_cell_mutation(cell: str) -> None:
    workbook, total_sheet, sources, groups = _written_formula_fixture()
    total_sheet[cell] = "MUTATED"
    with pytest.raises(ExcelExportError, match="mismatch"):
        _validate_written_global_formulas(
            workbook, total_sheet, sources, groups
        )


@pytest.mark.parametrize("cell", ["B7", "D7", "F7", "I7", "J7", "K7"])
def test_written_formula_validator_rejects_source_cell_mutation(
    cell: str,
) -> None:
    workbook, total_sheet, sources, groups = _written_formula_fixture()
    workbook["долна"][cell] = "MUTATED"
    with pytest.raises(ExcelExportError, match="source value/formula mismatch"):
        _validate_written_global_formulas(
            workbook, total_sheet, sources, groups
        )


def test_written_formula_validator_rejects_bad_source_token_or_sheet() -> None:
    workbook, total_sheet, sources, groups = _written_formula_fixture()
    workbook["долна"]["T7"] = "MUTATED"
    with pytest.raises(ExcelExportError, match="Source hidden"):
        _validate_written_global_formulas(
            workbook, total_sheet, sources, groups
        )

    workbook, total_sheet, sources, groups = _written_formula_fixture()
    key = groups[0].key
    source_group = sources[key][0][2]
    missing_sources = {key: [("липсва", 7, source_group)]}
    with pytest.raises(ExcelExportError, match="missing source sheet"):
        _validate_written_global_formulas(
            workbook, total_sheet, missing_sources, groups
        )


def test_sumif_range_filters_interleaved_a_b_a_identities() -> None:
    base = replace(
        _position(1, 1), master_topology_signature="sha256:base"
    )
    other = replace(base, diameter_mm=14)
    group_a1 = ExportGroup(key=base.group_key, positions=[base])
    group_b = ExportGroup(key=other.group_key, positions=[other])
    group_a2 = ExportGroup(
        key=base.group_key,
        positions=[replace(base, annotation_handle="A2")],
    )
    refs = [
        ("долна", 7, group_a1),
        ("долна", 9, group_a2),
    ]
    token_a = _production_identity_token(base.group_key)
    token_b = _production_identity_token(other.group_key)
    assert token_a != token_b
    assert _sumif_identity_formula(
        refs, total_row=7, value_column="B"
    ) == "=SUMIF('долна'!$T$7:$T$9,$T7,'долна'!$B$7:$B$9)"
    assert [token_a, token_b, token_a].count(token_a) == 2


def test_global_validator_rejects_structural_mutations(monkeypatch) -> None:
    import lion_rebar.export.excel_export as exporter

    position = replace(
        _position(1, 2), master_topology_signature="sha256:base"
    )
    source_group = ExportGroup(key=position.group_key, positions=[position])
    total_group = ExportGroup(key=position.group_key, positions=[position])
    sources = {position.group_key: [("долна", 7, source_group)]}
    numbers = {position.group_key: 1}
    assert _validate_global_aggregation(
        [position], sources, [total_group], numbers
    )["count_conserved"] is True

    with pytest.raises(ExcelExportError, match="duplicate"):
        _validate_global_aggregation(
            [position], sources, [total_group, total_group], numbers
        )
    with pytest.raises(ExcelExportError, match="key coverage"):
        _validate_global_aggregation([position], {}, [total_group], numbers)
    with pytest.raises(ExcelExportError, match="unique contiguous"):
        _validate_global_aggregation(
            [position], sources, [total_group], {position.group_key: 2}
        )
    wrong_count = ExportGroup(
        key=position.group_key,
        positions=[replace(position, count=3)],
    )
    with pytest.raises(ExcelExportError, match="count conservation"):
        _validate_global_aggregation(
            [position],
            {position.group_key: [("долна", 7, wrong_count)]},
            [total_group],
            numbers,
        )

    original_mass = exporter._group_formula_mass_kg
    monkeypatch.setattr(
        exporter,
        "_group_formula_mass_kg",
        lambda group: original_mass(group) + (1.0 if group is total_group else 0.0),
    )
    with pytest.raises(ExcelExportError, match="mass conservation"):
        _validate_global_aggregation(
            [position], sources, [total_group], numbers
        )


def test_total_formula_scales_past_255_same_identity_rows(
    tmp_path: Path, monkeypatch,
) -> None:
    import lion_rebar.export.excel_export as exporter
    import openpyxl.drawing.image as openpyxl_image
    import openpyxl.reader.excel as openpyxl_reader

    monkeypatch.setattr(openpyxl_image, "PILImage", False)
    monkeypatch.setattr(
        openpyxl_reader, "find_images", lambda *_args, **_kwargs: ([], [])
    )

    def fake_render(path: Path, **_kwargs) -> Path:
        path.write_bytes(base64.b64decode(_PNG_1X1))
        return path

    monkeypatch.setattr(exporter, "render_shape_png", fake_render)
    frame = _frame(1, "BEAMS", "греди")
    positions = [
        replace(
            _position(index, 1),
            frame_index=1,
            beam_name="Гр. STRESS 25/50",
            beam_label_handle="BEAM-STRESS",
        )
        for index in range(1, 301)
    ]
    output = tmp_path / "rc40_300_rows.xlsx"
    result = generate_excel_export(
        ProjectModel(name="RC40-STRESS", root=str(tmp_path)),
        _engine(),
        [frame],
        [*positions],
        [ExportFrameConfig(1, "BEAMS", "греди", frame.combined_title)],
        output,
    )

    workbook = load_workbook(output, data_only=False)
    for cell in ("B7", "J7", "K7"):
        formula = workbook["Общо"][cell].value
        assert formula.count("SUMIF(") == 1
        assert len(formula) < 256
    assert "$T$8:$T$307" in workbook["Общо"]["B7"].value
    total_token = workbook["Общо"]["T7"].value
    assert all(
        workbook["греди"][f"T{row}"].value == total_token
        for row in range(8, 308)
    )
    assert result.grouped_row_count == 300
    assert result.total_group_count == 1
    trace = json.loads(result.traceability_path.read_text(encoding="utf-8"))
    assert trace["global_aggregation"]["max_source_groups_per_identity"] == 300
    assert trace["global_aggregation"]["max_source_sheets_per_identity"] == 1
    assert trace["global_aggregation"]["source_identity_cells_validated"] == 300
    assert trace["global_aggregation"]["source_value_cells_validated"] == 1800
