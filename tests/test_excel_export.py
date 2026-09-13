from __future__ import annotations

import base64
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from lion_rebar.engine.models import EnginePackage, SegmentDefinition, ShapeDefinition
from lion_rebar.export.excel_export import (
    ExportFrameConfig,
    ExportPosition,
    _resolve_active_shape_code,
    generate_excel_export,
)
from lion_rebar.export.frame_detection import DrawingFrame, classify_title
from lion_rebar.project.models import ProjectModel


def _engine() -> EnginePackage:
    return EnginePackage(
        engine_id="test-engine",
        title="TEST",
        version="1.0",
        content_sha256="abc123",
        production_ready=True,
        shapes=[
            ShapeDefinition(code="F001", name="ПРАВ ПРЪТ", state="LOCKED", locked=True, revision=2, segments=[SegmentDefinition(index=1, length_code="L1")]),
            ShapeDefinition(code="F004", name="ПРЪТ С ЕДИН ВИНКЕЛ", state="LOCKED", locked=True, revision=3, segments=[SegmentDefinition(index=1, length_code="L1", bend_after_deg=90, bend_direction="RIGHT"), SegmentDefinition(index=2, length_code="L2")]),
        ],
    )


def test_title_classification_with_separate_reinforcement_text() -> None:
    assert classify_title("ДОЛНА АРМИРОВКА НА КОТА +9.50", "/УСИЛВАЩА МРЕЖА/")[0] == "BOTTOM_EXTRA"
    assert classify_title("ГОРНА АРМИРОВКА НА КОТА +9.50", "/ОСНОВНА МРЕЖА/")[0] == "TOP"
    assert classify_title("АРМИРОВКА ГРЕДИ +9.50", "")[0] == "BEAMS"


def test_retired_codes_follow_only_explicit_engine_replacements() -> None:
    active = ShapeDefinition(code="F010", name="CANONICAL", active=True)
    alias = ShapeDefinition(
        code="F009",
        name="ALIAS",
        active=False,
        state="RETIRED",
        alias_of="F010",
    )
    replacement = ShapeDefinition(
        code="F005",
        name="REPLACED",
        active=False,
        state="RETIRED",
        recognition={"canonical_replacement": "F003"},
    )
    u_shape = ShapeDefinition(code="F003", name="U", active=True)
    mapping = {shape.code: shape for shape in (active, alias, replacement, u_shape)}

    assert _resolve_active_shape_code("F009", all_shape_map=mapping) == "F010"
    assert _resolve_active_shape_code("F005", all_shape_map=mapping) == "F003"

    distinct_fiba = ShapeDefinition(code="F005", name="ФИБА", active=True, state="LOCKED")
    distinct_mapping = {"F003": u_shape, "F005": distinct_fiba}
    assert _resolve_active_shape_code("F005", all_shape_map=distinct_mapping) == "F005"


def test_dynamic_sheets_live_formulas_and_vid_fit(tmp_path: Path, monkeypatch) -> None:
    import openpyxl.drawing.image as openpyxl_image
    import lion_rebar.export.excel_export as exporter

    monkeypatch.setattr(openpyxl_image, "PILImage", False)

    def fake_render(path, **_kwargs):
        path.write_bytes(base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Z7r0AAAAASUVORK5CYII="
        ))
        return path
    monkeypatch.setattr(exporter, "render_shape_png", fake_render)
    project = ProjectModel(name="Test", root=str(tmp_path))
    frames = [
        DrawingFrame(1, "A", 0, 0, "ДОЛНА АРМИРОВКА", "/ОСНОВНА МРЕЖА/", "BOTTOM", "долна", "HIGH"),
        DrawingFrame(2, "B", 1000, 0, "АРМИРОВКА ГРЕДИ", "", "BEAMS", "греди", "HIGH"),
    ]
    positions = [
        ExportPosition("T1", "15N12x200", 15, 12, 200, "F001", "ПРАВ ПРЪТ", 2, ("G1",), 1, (200,), ()),
        ExportPosition("T2", "15N12x200", 15, 12, 200, "F001", "ПРАВ ПРЪТ", 2, ("G2",), 1, (200,), ()),
        ExportPosition("T3", "3N16x1150", 3, 16, 1150, "F004", "ПРЪТ С ЕДИН ВИНКЕЛ", 3, ("G3", "G4"), 2, (1125, 25), (-90,), beam_name="Гр. A.1 25/40", beam_label_handle="BL1"),
        ExportPosition("T4", "5N12x200", 5, 12, 200, "F001", "ПРАВ ПРЪТ", 2, ("G5",), 2, (200,), (), beam_name="Гр. A.2 25/50", beam_label_handle="BL2"),
    ]
    configs = [
        ExportFrameConfig(1, "BOTTOM", "долна", frames[0].combined_title),
        ExportFrameConfig(2, "BEAMS", "греди", frames[1].combined_title),
    ]
    output = tmp_path / "export.xlsx"
    result = generate_excel_export(project, _engine(), frames, positions, configs, output)
    assert result.safe_position_count == 4
    assert result.position_numbers_by_annotation == {
        "T1": 1,
        "T2": 1,
        "T3": 2,
        "T4": 1,
    }
    import openpyxl.reader.excel as openpyxl_reader
    monkeypatch.setattr(openpyxl_reader, "find_images", lambda *_args, **_kwargs: ([], []))
    wb = load_workbook(output, data_only=False)
    assert wb.sheetnames == ["долна", "греди", "Общо"]

    # One global position registry: F001/Ø12/L200 is position 1 everywhere.
    assert wb["долна"]["A7"].value == 1
    assert wb["долна"]["B7"].value == "=15+15"
    assert wb["греди"]["A8"].value == 2  # F004 first appears in beam A.1
    assert wb["греди"]["A10"].value == 1  # repeated F001 keeps position 1
    assert wb["Общо"]["A7"].value == 1
    assert wb["Общо"]["A8"].value == 2

    # Beam sheet is split into named beam sections and does not merge rows
    # across different beams.
    assert wb["греди"]["A7"].value.startswith("Гр. A.1")
    assert wb["греди"]["A9"].value.startswith("Гр. A.2")
    assert "A7:K7" in {str(rng) for rng in wb["греди"].merged_cells.ranges}
    assert "A9:K9" in {str(rng) for rng in wb["греди"].merged_cells.ranges}

    assert wb["долна"]["I7"].value == "=(D7/2)*(D7/2)*3.14*0.007859872"
    assert wb["долна"]["J7"].value == "=B7*F7/100"
    assert wb["долна"]["K7"].value == "=J7*I7"
    assert wb["Общо"]["B7"].value.startswith("=SUMIF(")
    assert "'долна'!$T$7:$T$7" in wb["Общо"]["B7"].value
    assert "'греди'!$T$10:$T$10" in wb["Общо"]["B7"].value
    assert wb["Общо"]["B7"].value.count("SUMIF(") == 2
    assert wb["Общо"]["J7"].value.count("SUMIF(") == 2
    assert wb["Общо"]["K7"].value.count("SUMIF(") == 2

    # Per-sheet kg totals and overall kg/tonnage remain live formulas.
    assert wb["долна"]["I8"].value == "ОБЩО КГ"
    assert wb["долна"]["K8"].value == "=SUM(K7:K7)"
    assert wb["греди"]["I11"].value == "ОБЩО КГ"
    assert wb["Общо"]["I9"].value == "ОБЩО КГ"
    assert wb["Общо"]["K9"].value == "=SUM(K7:K8)"
    assert wb["Общо"]["I10"].value == "ОБЩО Т"
    assert wb["Общо"]["K10"].value == "=K9/1000"

    assert wb["долна"].column_dimensions["G"].width >= 46
    assert wb["долна"].row_dimensions[7].height >= 112
    # Column G is a normal cell in the continuous A:K grid. Its four borders
    # remain visible around the anchored VID image.
    g_border = wb["долна"]["G7"].border
    assert g_border.left.style == "thin"
    assert g_border.right.style == "thin"
    assert g_border.top.style == "thin"
    assert g_border.bottom.style == "thin"
    with zipfile.ZipFile(output) as archive:
        media = [name for name in archive.namelist() if name.startswith("xl/media/") and name.endswith(".png")]
        drawings = [name for name in archive.namelist() if name.startswith("xl/drawings/drawing") and name.endswith(".xml")]
    assert len(media) == 5
    assert drawings
    assert wb["долна"].column_dimensions["L"].hidden is True
    assert wb["долна"].column_dimensions["T"].hidden is True
    assert result.traceability_path.is_file()


def test_master_v031_preflight_and_f006_renderer_are_locked() -> None:
    from lion_rebar.export.master_engine_v031 import (
        MASTER_ENGINE_SHA256,
        master_engine_preflight,
        parameters_from_segments,
        render_master_svg,
    )

    preflight = master_engine_preflight()
    assert preflight.sha256 == MASTER_ENGINE_SHA256
    assert preflight.form_codes == tuple(f"F{index:03d}" for index in range(1, 10))
    assert preflight.tests_passed == preflight.test_count == 10

    params = parameters_from_segments(
        "F006",
        (10, 45, 25, 45, 25, 10),
        production_length_cm=160,
    )
    assert params.values == {"L1": 10.0, "L2": 45.0, "L3": 25.0}
    svg = render_master_svg("F006", params.values)
    assert "45 cm" in svg
    assert "25 cm" in svg
    assert "10 cm" in svg
    assert '<rect class="rebar-line"' in svg
    # The locked F006 has one rectangular contour and two separately drawn hooks.
    assert svg.count('<line class="rebar-line"') == 2


def test_contextual_rule_keeps_225_plus_15_for_2n8x240() -> None:
    from types import SimpleNamespace

    from lion_rebar.export.excel_export import _segments_for_link

    link = {
        "annotation_handle": "C20D::BJ-FI-L",
        "annotation_text": "2N8x240",
        "candidate_notes": [
            "Приложено е одобрено проектно правило за средно желязо N8 в греда: "
            "225 + 15 cm, винкел 90°. Наклонът на късото рамо в DXF се приема за символично изчертаване."
        ],
    }
    segments, turns = _segments_for_link(
        link,
        shape=SimpleNamespace(code="F004", segments=[]),
        entity_map={},
        production_length_cm=240,
    )
    assert segments == (225.0, 15.0)
    assert turns == (90.0,)


def test_dimension_texts_override_schematic_dxf_for_9n10x215() -> None:
    from types import SimpleNamespace

    from lion_rebar.dxf.models import DxfEntityRecord
    from lion_rebar.export.excel_export import _segments_for_link

    entities = [
        DxfEntityRecord("L1", "LINE", "Arm - N10", {"start": [0, 20, 0], "end": [0, 0, 0]}),
        DxfEntityRecord("L2", "LINE", "Arm - N10", {"start": [0, 0, 0], "end": [176.6, 0, 0]}),
        DxfEntityRecord("L3", "LINE", "Arm - N10", {"start": [176.6, 0, 0], "end": [176.6, 25, 0]}),
        DxfEntityRecord("D1", "TEXT", "koti", {"insert": [-8, 10, 0], "rotation": 90}, text="20"),
        DxfEntityRecord("D2", "TEXT", "koti", {"insert": [88, -8, 0], "rotation": 0}, text="175"),
        DxfEntityRecord("D3", "TEXT", "koti", {"insert": [184, 12, 0], "rotation": 90}, text="20"),
    ]
    entity_map = {record.handle: record for record in entities}
    link = {
        "annotation_handle": "8F70::BJ-FI-L",
        "annotation_text": "9N10x215",
        "geometry_handles": ["L1", "L2", "L3"],
    }
    segments, _turns = _segments_for_link(
        link,
        shape=SimpleNamespace(code="F003", segments=[]),
        entity_map=entity_map,
        production_length_cm=215,
    )
    assert segments == (20.0, 175.0, 20.0)


def test_persisted_auto_f003_dimensions_survive_export_resolution() -> None:
    from types import SimpleNamespace

    from lion_rebar.dxf.models import DxfEntityRecord
    from lion_rebar.export.excel_export import _segments_for_link

    # The single schematic line cannot prove an F003. Export must consume the
    # vector already validated and frozen by the AUTO batch gate.
    geometry = DxfEntityRecord(
        "G1",
        "LINE",
        "Arm - N8",
        {"start": [0, 0, 0], "end": [100, 0, 0]},
    )
    segments, turns = _segments_for_link(
        {
            "annotation_handle": "55E5::BJ-FI-L",
            "annotation_text": "40N8x740",
            "geometry_handles": ["G1"],
            "production_segment_lengths_cm": [15.0, 710.0, 15.0],
            "production_turn_angles_deg": [90.0, 90.0],
        },
        shape=SimpleNamespace(code="F003", segments=[]),
        entity_map={"G1": geometry},
        production_length_cm=740.0,
    )
    assert segments == (15.0, 710.0, 15.0)
    assert turns == (90.0, 90.0)


def test_f006_dimension_texts_replace_schematic_contour_lengths() -> None:
    from types import SimpleNamespace

    from lion_rebar.dxf.models import DxfEntityRecord
    from lion_rebar.export.excel_export import _segments_for_link

    # The DXF is schematic (64x34 with 20 cm graphic keys), while the locked
    # MASTER dimensions printed beside it are 45x25 with 10 cm hooks.
    entities = [
        DxfEntityRecord("R1", "LINE", "ARM_Bars", {"start": [0, 0, 0], "end": [64, 0, 0]}),
        DxfEntityRecord("R2", "LINE", "ARM_Bars", {"start": [64, 0, 0], "end": [64, 34, 0]}),
        DxfEntityRecord("R3", "LINE", "ARM_Bars", {"start": [64, 34, 0], "end": [0, 34, 0]}),
        DxfEntityRecord("R4", "LINE", "ARM_Bars", {"start": [0, 34, 0], "end": [0, 0, 0]}),
        DxfEntityRecord("H1", "LINE", "ARM_Bars", {"start": [50, 34, 0], "end": [64, 48, 0]}),
        DxfEntityRecord("H2", "LINE", "ARM_Bars", {"start": [64, 25, 0], "end": [50, 11, 0]}),
        DxfEntityRecord("DW", "TEXT", "koti", {"insert": [32, -8, 0], "rotation": 0}, text="45"),
        DxfEntityRecord("DH", "TEXT", "koti", {"insert": [-8, 17, 0], "rotation": 90}, text="25"),
        DxfEntityRecord("DK", "TEXT", "koti", {"insert": [68, 41, 0], "rotation": 45}, text="10"),
    ]
    entity_map = {record.handle: record for record in entities}
    link = {
        "annotation_handle": "S1::BJ-FI-L",
        "annotation_text": "11N8x160",
        "geometry_handles": ["R1", "R2", "R3", "R4", "H1", "H2"],
        # Historical payload order must not bypass the explicit dimensions.
        "canonical_segment_lengths_cm": [64, 34, 64, 34, 20, 20],
    }
    segments, _turns = _segments_for_link(
        link,
        shape=SimpleNamespace(code="F006", segments=[]),
        entity_map=entity_map,
        production_length_cm=160,
    )
    assert segments == (10.0, 45.0, 25.0, 45.0, 25.0, 10.0)


def test_master_svg_rasterizes_in_runtime(tmp_path: Path) -> None:
    """Run Qt SVG rasterization in an isolated process with a hard timeout.

    A renderer regression must fail with a useful test result instead of
    freezing the complete managed update for five minutes.
    """
    import os
    import subprocess
    import sys

    import pytest

    pytest.importorskip("PySide6")
    output = tmp_path / "f006.png"
    project_root = Path(__file__).resolve().parents[1]
    script = f"""
import faulthandler
from pathlib import Path
faulthandler.dump_traceback_later(20, exit=True)
from lion_rebar.engine.models import ShapeDefinition
from lion_rebar.export.shape_image import render_shape_png
shape = ShapeDefinition(code='F006', name='СТРЕМЕ (БИГЛА)', state='LOCKED', locked=True)
output = Path({str(output)!r})
render_shape_png(
    output,
    shape=shape,
    segment_lengths_cm=(10, 45, 25, 45, 25, 10),
    production_length_cm=160,
)
print(output.stat().st_size)
"""
    env = os.environ.copy()
    env.setdefault("QT_QPA_PLATFORM", "offscreen")
    existing_pythonpath = env.get("PYTHONPATH", "")
    env["PYTHONPATH"] = str(project_root) + (os.pathsep + existing_pythonpath if existing_pythonpath else "")
    try:
        completed = subprocess.run(
            [sys.executable, "-c", script],
            cwd=project_root,
            env=env,
            capture_output=True,
            text=True,
            timeout=35,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise AssertionError(
            "MASTER SVG rasterization exceeded 35 seconds. "
            f"stdout={exc.stdout!r}; stderr={exc.stderr!r}"
        ) from exc

    assert completed.returncode == 0, completed.stdout + "\n" + completed.stderr
    data = output.read_bytes()
    assert data.startswith(b"\x89PNG\r\n\x1a\n")
    assert len(data) > 1000

def test_master_svg_alpha_crop_avoids_qbitmap_region_path() -> None:
    import sys

    from lion_rebar.export.shape_image import _visible_alpha_bounds

    class FakeRect:
        def __init__(self, x=0, y=0, width=0, height=0):
            self.values = (x, y, width, height)

    class FakeArgb32Image:
        def __init__(self, width: int, height: int):
            self._width = width
            self._height = height
            self._bytes_per_line = width * 4
            self.buffer = bytearray(self._bytes_per_line * height)

        def width(self):
            return self._width

        def height(self):
            return self._height

        def bytesPerLine(self):
            return self._bytes_per_line

        def constBits(self):
            return self.buffer

        def set_alpha(self, x: int, y: int, alpha: int):
            offset = 3 if sys.byteorder == "little" else 0
            self.buffer[y * self._bytes_per_line + x * 4 + offset] = alpha

    image = FakeArgb32Image(20, 16)
    for y in range(5, 11):
        for x in range(4, 14):
            image.set_alpha(x, y, 255)

    bounds = _visible_alpha_bounds(image, FakeRect)
    assert bounds.values == (4, 5, 10, 6)


def test_f006_master_orientation_normalizes_swapped_rectangle_sides() -> None:
    from lion_rebar.export.master_engine_v031 import parameters_from_segments, render_master_svg

    params = parameters_from_segments(
        "F006",
        (10, 25, 45, 25, 45, 10),
        production_length_cm=160,
    )
    assert params.values == {"L1": 10.0, "L2": 45.0, "L3": 25.0}
    assert params.expanded_segments_cm == (10.0, 45.0, 25.0, 45.0, 25.0, 10.0)
    svg = render_master_svg("F006", params.values)
    assert "45 cm" in svg
    assert "25 cm" in svg


def test_master_export_svg_uses_readable_dimensions_and_finer_rebar_lines() -> None:
    from lion_rebar.export.master_engine_v031 import render_master_svg

    svg = render_master_svg("F004", {"L1": 225, "L2": 15})
    assert "stroke-width:1.2" in svg
    assert "font-size:4.6px" in svg
    assert "font-family:Arial,sans-serif" in svg
    assert 'viewBox="-12 -10 124 120"' in svg
    assert "225 cm" in svg
    assert "15 cm" in svg


def test_beam_sheet_keeps_duplicate_identity_annotations_as_separate_rows_but_total_aggregates(
    tmp_path: Path,
    monkeypatch,
) -> None:
    import openpyxl.drawing.image as openpyxl_image
    import lion_rebar.export.excel_export as exporter

    monkeypatch.setattr(openpyxl_image, "PILImage", False)

    def fake_render(path, **_kwargs):
        path.write_bytes(base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Z7r0AAAAASUVORK5CYII="
        ))
        return path

    monkeypatch.setattr(exporter, "render_shape_png", fake_render)
    project = ProjectModel(name="Beam rows", root=str(tmp_path))
    frame = DrawingFrame(
        1,
        "B",
        0,
        0,
        "АРМИРОВКА ГРЕДИ",
        "",
        "BEAMS",
        "греди",
        "HIGH",
    )
    positions = [
        ExportPosition(
            "T1",
            "2N12x200",
            2,
            12,
            200,
            "F001",
            "ПРАВ ПРЪТ",
            2,
            ("G1",),
            1,
            (200,),
            (),
            beam_name="Гр. B.1 25/40",
            beam_label_handle="BL1",
        ),
        ExportPosition(
            "T2",
            "3N12x200",
            3,
            12,
            200,
            "F001",
            "ПРАВ ПРЪТ",
            2,
            ("G2",),
            1,
            (200,),
            (),
            beam_name="Гр. B.1 25/40",
            beam_label_handle="BL1",
        ),
    ]
    output = tmp_path / "beam_rows.xlsx"
    result = generate_excel_export(
        project,
        _engine(),
        [frame],
        positions,
        [ExportFrameConfig(1, "BEAMS", "греди", frame.combined_title)],
        output,
    )

    import openpyxl.reader.excel as openpyxl_reader

    monkeypatch.setattr(openpyxl_reader, "find_images", lambda *_args, **_kwargs: ([], []))
    wb = load_workbook(output, data_only=False)
    assert wb.sheetnames == ["греди", "Общо"]
    assert wb["греди"]["A7"].value == "Гр. B.1 25/40"
    assert wb["греди"]["A8"].value == wb["греди"]["A9"].value == 1
    assert wb["греди"]["B8"].value == "=2"
    assert wb["греди"]["B9"].value == "=3"
    assert wb["греди"]["O8"].value == "T1"
    assert wb["греди"]["O9"].value == "T2"
    assert wb["Общо"]["A7"].value == 1
    assert wb["Общо"]["B7"].value == (
        "=SUMIF('греди'!$T$8:$T$9,$T7,'греди'!$B$8:$B$9)"
    )
    assert result.grouped_row_count == 2
    assert result.total_group_count == 1
