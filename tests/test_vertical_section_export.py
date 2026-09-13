from __future__ import annotations

import base64
import csv
import json
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from lion_rebar.dxf.models import DxfEntityRecord
from lion_rebar.engine.models import EnginePackage, SegmentDefinition, ShapeDefinition
from lion_rebar.export.errors import ExcelExportError
from lion_rebar.export.excel_export import (
    ExportFrameConfig,
    ExportGroup,
    ExportPosition,
    VerticalSectionLabel,
    _assign_vertical_section,
    _detect_vertical_section_labels,
    _group_positions,
    _validate_source_coverage,
    _validate_vertical_assignments,
    _vertical_sections,
    generate_excel_export,
)
from lion_rebar.export.frame_detection import DrawingFrame
from lion_rebar.project.models import ProjectModel
from lion_rebar.project.data_identity import DataIdentity
from lion_rebar.project.integrity import _freeze_beam_assignments, _write_ledger
from lion_rebar.review_center_model import ACTIONABLE_CLASS


def _frame() -> DrawingFrame:
    return DrawingFrame(
        1,
        "FRAME",
        0.0,
        0.0,
        "VERTICAL REINFORCEMENT",
        "",
        "VERTICALS",
        "verticals",
        "HIGH",
    )


def _type_title(index: int) -> str:
    return (
        f"\u0422\u0438\u043f {index} - 100x25 - 1 "
        "\u0431\u0440.(\u0421\u0445.\u0430\u0440\u043c.1 )"
    )


def _title_entity(handle: str, text: str, y: float) -> DxfEntityRecord:
    return DxfEntityRecord(
        handle,
        "TEXT",
        "DETAIL",
        {
            "insert": [0.0, y, 0.0],
            "visual_anchor": [250.0, y + 15.0, 0.0],
            "visual_bbox": [
                [0.0, y, 0.0],
                [500.0, y, 0.0],
                [0.0, y + 30.0, 0.0],
                [500.0, y + 30.0, 0.0],
            ],
        },
        text,
    )


def _divider_entity(handle: str, y: float) -> DxfEntityRecord:
    return DxfEntityRecord(
        handle,
        "LINE",
        "DETAIL",
        {"start": [0.0, y, 0.0], "end": [500.0, y, 0.0]},
    )


def _synthetic_labels() -> list[VerticalSectionLabel]:
    entities: list[DxfEntityRecord] = []
    for index in range(17):
        title_y = 18_000.0 - index * 1_000.0
        entities.append(
            _title_entity(f"TITLE_{index + 1}", _type_title(index + 1), title_y)
        )
        entities.append(_divider_entity(f"DIVIDER_{index + 1}", title_y - 20.0))
    # A generic plan label is deliberately ignored: no full schema/count text.
    entities.append(
        _title_entity(
            "PLAN_LABEL",
            "\u0422\u0438\u043f 1",
            500.0,
        )
    )
    return _detect_vertical_section_labels(entities, [_frame()])


def _core_title(index: int, axis: str = "A") -> str:
    return (
        f"\u0421\u0442\u0431.\u042f\u0434\u0440\u043e {axis}.{index} - "
        "(\u0421\u0445.\u0430\u0440\u043c.2 )"
    )


def _named_title(identifier: str, dimensions: str) -> str:
    return (
        f"{identifier} {dimensions} - 1 "
        "\u0431\u0440.(\u0421\u0445.\u0430\u0440\u043c.1 )"
    )


def test_spatially_interleaved_real_titles_are_output_in_semantic_blocks() -> None:
    # This is the real drawing's interleaving pattern: A2 is spatially below
    # Type 1/1b, and the named verticals are not in natural title order.
    spatial_titles = [
        ("CORE_A1", _core_title(1)),
        ("CORE_A3", _core_title(3)),
        ("TYPE_1", _type_title(1)),
        (
            "TYPE_1B",
            "\u0422\u0438\u043f 1b - 25x25 - 10 "
            "\u0431\u0440.(\u0421\u0445.\u0430\u0440\u043c.1 )",
        ),
        ("CORE_A2", _core_title(2)),
        (
            "TYPE_1A",
            "\u0422\u0438\u043f 1a - 170x25 - 2 "
            "\u0431\u0440.(\u0421\u0445.\u0430\u0440\u043c.2 )",
        ),
        (
            "TYPE_1C",
            "\u0422\u0438\u043f 1c - 150x25 - 2 "
            "\u0431\u0440.(\u0421\u0445.\u0430\u0440\u043c.1 )",
        ),
        ("TYPE_2", _type_title(2)),
        (
            "TYPE_2A",
            "\u0422\u0438\u043f 2a - 350x25 - 3 "
            "\u0431\u0440.(\u0421\u0445.\u0430\u0440\u043c.2 )",
        ),
        (
            "TYPE_2B",
            "\u0422\u0438\u043f 2b - 250x25 - 2 "
            "\u0431\u0440.(\u0421\u0445.\u0430\u0440\u043c.2 )",
        ),
        (
            "TYPE_2C",
            "\u0422\u0438\u043f 2c - 215x25 - 3 "
            "\u0431\u0440.(\u0421\u0445.\u0430\u0440\u043c.2 )",
        ),
        (
            "NAMED_SA20",
            _named_title("\u0428.A-20-1", "197x25"),
        ),
        ("NAMED_KA76", _named_title("K.A-76", "220x25")),
        (
            "NAMED_KSA16",
            _named_title("K.\u0428.A16", ""),
        ),
        ("NAMED_KA40", _named_title("K.A.40", "")),
        ("NAMED_KA21", _named_title("K.A.21", "60x40")),
        ("NAMED_KA11", _named_title("K.A.11,11a", "88x25")),
    ]
    entities: list[DxfEntityRecord] = []
    for spatial_index, (handle, title) in enumerate(spatial_titles):
        title_y = 18_000.0 - spatial_index * 1_000.0
        entities.append(_title_entity(handle, title, title_y))
        entities.append(_divider_entity(f"D_{handle}", title_y - 20.0))

    labels = _detect_vertical_section_labels(entities, [_frame()])
    assert [label.title_handle for label in labels] == [
        "CORE_A1",
        "CORE_A2",
        "CORE_A3",
        "TYPE_1",
        "TYPE_1A",
        "TYPE_1B",
        "TYPE_1C",
        "TYPE_2",
        "TYPE_2A",
        "TYPE_2B",
        "TYPE_2C",
        "NAMED_KA11",
        "NAMED_KA21",
        "NAMED_KA40",
        "NAMED_KA76",
        "NAMED_KSA16",
        "NAMED_SA20",
    ]
    assert [label.kind for label in labels] == (
        ["CORE"] * 3
        + ["TYPE"] * 8
        + ["NAMED_VERTICAL"] * 6
    )
    assert [label.order for label in labels] == list(range(1, 18))


def test_b_axis_sections_use_the_same_strict_title_divider_contract() -> None:
    spatial_titles = [
        ("CORE_B1", _core_title(1, "B")),
        ("CORE_CYR_B2", _core_title(2, "\u0412")),
        ("TYPE_1", _type_title(1)),
        ("NAMED_KB11", _named_title("K.B.11,11a", "88x25")),
        ("NAMED_KSHB16", _named_title("K.\u0428.B16", "60x40")),
        ("NAMED_CYR_B5", _named_title("\u0428.\u0412.5", "277x25")),
    ]
    entities: list[DxfEntityRecord] = []
    divider_y_by_title: dict[str, float] = {}
    for spatial_index, (handle, title) in enumerate(spatial_titles):
        title_y = 6_000.0 - spatial_index * 1_000.0
        divider_y = title_y - 20.0
        divider_y_by_title[handle] = divider_y
        entities.append(_title_entity(handle, title, title_y))
        entities.append(_divider_entity(f"D_{handle}", divider_y))

    # A visual plan label is not section evidence without the complete
    # Schematic reinforcement production title. It must remain ignored.
    entities.append(
        _title_entity(
            "PLAN_CORE_B",
            "\u0421\u0442\u0431.\u042f\u0434\u0440\u043e B.9",
            500.0,
        )
    )

    labels = _detect_vertical_section_labels(entities, [_frame()])
    assert [label.title_handle for label in labels] == [
        "CORE_B1",
        "CORE_CYR_B2",
        "TYPE_1",
        "NAMED_KB11",
        "NAMED_KSHB16",
        "NAMED_CYR_B5",
    ]
    assert [label.kind for label in labels] == [
        "CORE",
        "CORE",
        "TYPE",
        "NAMED_VERTICAL",
        "NAMED_VERTICAL",
        "NAMED_VERTICAL",
    ]
    assert [label.order for label in labels] == list(range(1, 7))

    labels_by_handle = {label.title_handle: label for label in labels}
    for handle, _title in spatial_titles:
        assignment = _assign_vertical_section(
            250.0,
            divider_y_by_title[handle] - 100.0,
            1,
            labels,
        )
        assert assignment is not None
        assert assignment.section == labels_by_handle[handle]


def test_b_axis_section_still_requires_one_explicit_divider() -> None:
    title = _title_entity("CORE_B1", _core_title(1, "B"), 1_000.0)
    with pytest.raises(ExcelExportError, match="no explicit divider"):
        _detect_vertical_section_labels([title], [_frame()])

    with pytest.raises(ExcelExportError, match="duplicate divider"):
        _detect_vertical_section_labels(
            [
                title,
                _divider_entity("D1", 980.0),
                _divider_entity("D2", 980.0),
            ],
            [_frame()],
        )


def _vertical_position(
    *,
    handle: str,
    length_cm: float,
    section: VerticalSectionLabel,
    assignment_score: float,
    assignment_margin: float,
    count: int = 1,
    shape_code: str = "F001",
    shape_revision: int = 1,
    segment_lengths_cm: tuple[float, ...] | None = None,
) -> ExportPosition:
    return ExportPosition(
        annotation_handle=handle,
        annotation_text=f"{count}N8x{int(length_cm)}",
        count=count,
        diameter_mm=8.0,
        production_length_cm=length_cm,
        shape_code=shape_code,
        shape_name=shape_code,
        shape_revision=shape_revision,
        geometry_handles=(),
        frame_index=1,
        segment_lengths_cm=segment_lengths_cm or (length_cm,),
        vertical_section_title=section.title,
        vertical_section_label_handle=section.title_handle,
        vertical_section_divider_handle=section.divider_handle,
        vertical_section_kind=section.kind,
        vertical_section_order=section.order,
        vertical_assignment_method="AUTO_VERTICAL_SECTION_V1",
        vertical_assignment_score=assignment_score,
        vertical_assignment_margin=assignment_margin,
    )


def test_159_vertical_annotations_build_114_section_groups_and_38_totals() -> None:
    labels = _synthetic_labels()
    assert len(labels) == 17
    assert [label.order for label in labels] == list(range(1, 18))
    assert {label.kind for label in labels} == {"TYPE"}

    expected_counts = [
        30, 31, 33, 3, 5, 2, 4, 6, 6, 6, 5, 4, 3, 5, 5, 6, 5
    ]
    expected_source_groups = [
        17, 17, 18, 3, 5, 2, 4, 5, 5, 5, 5, 4, 3, 5, 5, 6, 5
    ]
    identity_sets = [
        list(range(0, 17)),
        list(range(17, 34)),
        list(range(20, 38)),
    ] + [
        list(range(group_count))
        for group_count in expected_source_groups[3:]
    ]
    assert sum(expected_counts) == 159
    assert sum(expected_source_groups) == 114
    positions: list[ExportPosition] = []
    assigned_counts: Counter[str] = Counter()
    serial = 0
    for section, section_count, identity_ids in zip(
        labels,
        expected_counts,
        identity_sets,
    ):
        for item_index in range(section_count):
            assignment = _assign_vertical_section(
                250.0,
                section.divider_y - 100.0 - item_index * 3.0,
                1,
                labels,
            )
            assert assignment is not None
            assert assignment.section.title_handle == section.title_handle
            serial += 1
            assigned_counts[section.title_handle] += 1
            identity_index = identity_ids[item_index % len(identity_ids)]
            length_cm = float(100 + identity_index)
            positions.append(
                _vertical_position(
                    handle=f"A{serial:03d}",
                    length_cm=length_cm,
                    section=assignment.section,
                    assignment_score=assignment.score,
                    assignment_margin=assignment.margin,
                )
            )

    assert len(positions) == 159
    assert list(assigned_counts.values()) == expected_counts
    _validate_vertical_assignments([_frame()], labels, positions)

    source_sections = _vertical_sections(positions)
    assert len(source_sections) == 17
    assert [len(section.groups) for section in source_sections] == expected_source_groups
    assert sum(len(section.groups) for section in source_sections) == 114
    assert sum(
        len(group.positions)
        for section in source_sections
        for group in section.groups
    ) == 159
    assert any(
        len(group.positions) > 1
        for section in source_sections
        for group in section.groups
    )

    total_groups = _group_positions(positions)
    assert len(total_groups) == 38
    assert len(total_groups) == len({position.group_key for position in positions})
    assert sum(len(group.positions) for group in total_groups) == 159


def test_vertical_grouping_uses_full_identity_and_never_crosses_sections() -> None:
    labels = [
        VerticalSectionLabel(1, "T1", "D1", "Type 1", "TYPE", 1, 0.0, 500.0, 1_000.0),
        VerticalSectionLabel(1, "T2", "D2", "Type 2", "TYPE", 2, 0.0, 500.0, 500.0),
    ]
    common = dict(
        length_cm=100.0,
        assignment_score=100.0,
        assignment_margin=400.0,
        shape_code="F003",
    )
    positions = [
        _vertical_position(
            handle="A1",
            section=labels[0],
            count=25,
            segment_lengths_cm=(10.0, 80.0, 10.0),
            **common,
        ),
        _vertical_position(
            handle="A2",
            section=labels[0],
            count=25,
            segment_lengths_cm=(10.0, 80.0, 10.0),
            **common,
        ),
        _vertical_position(
            handle="A3",
            section=labels[1],
            count=50,
            segment_lengths_cm=(10.0, 80.0, 10.0),
            **common,
        ),
        _vertical_position(
            handle="A4",
            section=labels[0],
            count=10,
            segment_lengths_cm=(20.0, 60.0, 20.0),
            **common,
        ),
        _vertical_position(
            handle="A5",
            section=labels[0],
            count=10,
            shape_revision=2,
            segment_lengths_cm=(10.0, 80.0, 10.0),
            **common,
        ),
    ]

    sections = _vertical_sections(positions)
    assert [len(section.groups) for section in sections] == [3, 1]
    merged = next(
        group
        for group in sections[0].groups
        if [position.annotation_handle for position in group.positions] == ["A1", "A2"]
    )
    assert merged.counts == [25, 25]
    assert all(
        {
            (
                position.frame_index,
                position.vertical_section_label_handle,
                position.vertical_section_divider_handle,
            )
            for position in group.positions
        }
        == {
            (
                group.representative.frame_index,
                group.representative.vertical_section_label_handle,
                group.representative.vertical_section_divider_handle,
            )
        }
        for section in sections
        for group in section.groups
    )


def test_source_coverage_validation_fails_closed_on_missing_or_duplicate_refs() -> None:
    section = VerticalSectionLabel(
        1, "T1", "D1", "Type 1", "TYPE", 1, 0.0, 500.0, 1_000.0
    )
    positions = [
        _vertical_position(
            handle=handle,
            length_cm=100.0,
            section=section,
            assignment_score=100.0,
            assignment_margin=400.0,
        )
        for handle in ("A1", "A2")
    ]
    group = ExportGroup(key=positions[0].group_key, positions=positions)
    with pytest.raises(ExcelExportError, match="coverage mismatch"):
        _validate_source_coverage(
            positions,
            {"A1": {"source_group_id": "verticals!B8"}},
            {group.key: [("verticals", 8, group)]},
            1,
        )

    locations = {
        handle: {"source_group_id": "verticals!B8"}
        for handle in ("A1", "A2")
    }
    with pytest.raises(ExcelExportError, match="exactly once"):
        _validate_source_coverage(
            positions,
            locations,
            {
                group.key: [
                    ("verticals", 8, group),
                    ("verticals", 9, group),
                ]
            },
            2,
        )


def test_vertical_assignment_and_divider_detection_fail_closed() -> None:
    labels = [
        VerticalSectionLabel(1, "T1", "D1", "S1", "TYPE", 1, 0.0, 100.0, 500.0),
        VerticalSectionLabel(1, "T2", "D2", "S2", "TYPE", 2, 0.0, 100.0, 500.0),
    ]
    assert _assign_vertical_section(50.0, 400.0, 1, labels) is None
    assert _assign_vertical_section(500.0, 400.0, 1, labels) is None

    title = _title_entity("TITLE", _type_title(1), 1_000.0)
    with pytest.raises(ExcelExportError, match="no explicit divider"):
        _detect_vertical_section_labels([title], [_frame()])

    duplicate_dividers = [
        title,
        _divider_entity("D1", 980.0),
        _divider_entity("D2", 980.0),
    ]
    with pytest.raises(ExcelExportError, match="duplicate divider"):
        _detect_vertical_section_labels(duplicate_dividers, [_frame()])


def test_final_control_freezes_vertical_fields_and_writes_them_to_ledger(
    tmp_path: Path,
) -> None:
    entities = [
        DxfEntityRecord(
            "FRAME",
            "TEXT",
            "$Ramka",
            {"insert": [0.0, 0.0, 0.0]},
            "\u0427\u0415\u0420\u0422\u0415\u0416:",
        ),
        DxfEntityRecord(
            "FRAME_TITLE",
            "TEXT",
            "TXT",
            {"insert": [100.0, 0.0, 0.0]},
            (
                "\u0410\u0420\u041c\u0418\u0420\u041e\u0412\u041a\u0410 "
                "\u0412\u0415\u0420\u0422\u0418\u041a\u0410\u041b\u0418"
            ),
        ),
        _title_entity("SECTION_TITLE", _type_title(1), 1_000.0),
        _divider_entity("SECTION_DIVIDER", 980.0),
        DxfEntityRecord(
            "A1",
            "BLOCK_ATTR_TEXT",
            "ARM_TXT",
            {
                "insert": [250.0, 800.0, 0.0],
                "visual_anchor": [250.0, 800.0, 0.0],
            },
            "1N8x100",
        ),
    ]
    extracted_payload = {
        "entities": [
            {
                "handle": entity.handle,
                "entity_type": entity.entity_type,
                "layer": entity.layer,
                "geometry": entity.geometry,
                "text": entity.text,
                "block_name": entity.block_name,
                "color_index": entity.color_index,
                "linetype": entity.linetype,
                "raw": entity.raw,
            }
            for entity in entities
        ]
    }
    rows = [
        {
            "annotation_class": ACTIONABLE_CLASS,
            "annotation_handle": "A1",
            "annotation_text": "1N8x100",
        }
    ]
    links = {
        "A1": {
            "annotation_handle": "A1",
            "annotation_text": "1N8x100",
            "resolved_annotation_text": "1N8x100",
            "production_count": 1,
            "production_diameter_mm": 8,
            "production_length_cm": 100,
            "master_shape_code": "F001",
            "master_shape_name": "STRAIGHT",
            "geometry_handles": [],
        }
    }
    identity = DataIdentity(
        project_id="P",
        source_sha256="S",
        source_generation_id="G",
        analysis_run_id="R",
    )
    frozen, issues = _freeze_beam_assignments(
        rows=rows,
        links_by_annotation=links,
        extracted_payload=extracted_payload,
        project_root=tmp_path,
        identity=identity,
    )
    assert issues == []
    assert frozen["A1"]["vertical_assignment_frozen"] is True
    assert frozen["A1"]["vertical_section_label_handle"] == "SECTION_TITLE"
    assert frozen["A1"]["vertical_section_divider_handle"] == "SECTION_DIVIDER"
    assert frozen["A1"]["vertical_assignment_method"] == "AUTO_VERTICAL_SECTION_V1"

    ledger_path = tmp_path / "final_integrity_ledger.csv"
    _write_ledger(
        path=ledger_path,
        rows=rows,
        links_by_annotation=links,
        decisions={},
        identity=identity,
        beam_assignments=frozen,
    )
    with ledger_path.open(encoding="utf-8-sig", newline="") as stream:
        ledger_row = next(csv.DictReader(stream))
    assert ledger_row["vertical_assignment_frozen"] == "1"
    assert ledger_row["vertical_section_title"] == _type_title(1)
    assert ledger_row["vertical_section_label_handle"] == "SECTION_TITLE"
    assert ledger_row["vertical_section_divider_handle"] == "SECTION_DIVIDER"
    assert ledger_row["vertical_section_order"] == "1"
    assert ledger_row["vertical_assignment_method"] == "AUTO_VERTICAL_SECTION_V1"


def _engine() -> EnginePackage:
    return EnginePackage(
        engine_id="vertical-test",
        title="VERTICAL TEST",
        version="1.0",
        content_sha256="vertical-test",
        production_ready=True,
        shapes=[
            ShapeDefinition(
                code="F001",
                name="STRAIGHT",
                state="LOCKED",
                locked=True,
                active=True,
                revision=1,
                segments=[SegmentDefinition(index=1, length_code="L1")],
            )
        ],
    )


def test_vertical_workbook_aggregates_within_section_and_keeps_full_trace(
    tmp_path: Path,
    monkeypatch,
) -> None:
    import openpyxl.drawing.image as openpyxl_image
    import lion_rebar.export.excel_export as exporter

    monkeypatch.setattr(openpyxl_image, "PILImage", False)

    def fake_render(path, **_kwargs):
        path.write_bytes(
            base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0l"
                "EQVR42mP8/x8AAusB9Y9Z7r0AAAAASUVORK5CYII="
            )
        )
        return path

    monkeypatch.setattr(exporter, "render_shape_png", fake_render)
    labels = [
        VerticalSectionLabel(1, "T1", "D1", "Type 1", "TYPE", 1, 0.0, 500.0, 1_000.0),
        VerticalSectionLabel(1, "T2", "D2", "Type 2", "TYPE", 2, 0.0, 500.0, 500.0),
    ]
    positions = [
        _vertical_position(
            handle="A1",
            length_cm=200.0,
            section=labels[0],
            assignment_score=100.0,
            assignment_margin=400.0,
            count=25,
        ),
        _vertical_position(
            handle="A2",
            length_cm=200.0,
            section=labels[0],
            assignment_score=120.0,
            assignment_margin=380.0,
            count=25,
        ),
        _vertical_position(
            handle="A3",
            length_cm=200.0,
            section=labels[1],
            assignment_score=100.0,
            assignment_margin=1_000_000_000.0,
            count=50,
        ),
    ]
    frame = _frame()
    output = tmp_path / "vertical_rows.xlsx"
    result = generate_excel_export(
        ProjectModel(name="Vertical rows", root=str(tmp_path)),
        _engine(),
        [frame],
        positions,
        [
            ExportFrameConfig(
                1,
                "VERTICALS",
                "verticals",
                frame.combined_title,
            )
        ],
        output,
    )

    import openpyxl.reader.excel as openpyxl_reader

    monkeypatch.setattr(
        openpyxl_reader,
        "find_images",
        lambda *_args, **_kwargs: ([], []),
    )
    workbook = load_workbook(output, data_only=False)
    source = workbook["verticals"]
    total = workbook[workbook.sheetnames[1]]
    assert source["A7"].value == "Type 1"
    assert source["A8"].value == 1
    assert source["B8"].value == "=25+25"
    assert source["O8"].value == "A1;A2"
    assert source["A9"].value == "Type 2"
    assert source["A10"].value == 1
    assert source["B10"].value == "=50"
    assert source["O10"].value == "A3"
    assert total["A7"].value == 1
    assert total["B7"].value == (
        "=SUMIF('verticals'!$T$8:$T$10,$T7,'verticals'!$B$8:$B$10)"
    )
    assert total["J7"].value == (
        "=SUMIF('verticals'!$T$8:$T$10,$T7,'verticals'!$J$8:$J$10)"
    )
    assert total["K7"].value == (
        "=SUMIF('verticals'!$T$8:$T$10,$T7,'verticals'!$K$8:$K$10)"
    )
    assert result.safe_position_count == 3
    assert result.grouped_row_count == 2
    assert result.total_group_count == 1

    trace = json.loads(result.traceability_path.read_text(encoding="utf-8"))
    assert trace["schema"] == "lion-rebar-excel-traceability/v3"
    assert trace["vertical_export"]["section_count"] == 2
    assert trace["vertical_export"]["approved_annotation_count"] == 3
    assert trace["vertical_export"]["source_group_row_count"] == 2
    assert trace["vertical_export"]["collapsed_annotation_count"] == 1
    assert trace["vertical_export"]["vertical_position_rows"] == 2
    assert trace["vertical_export"]["one_row_per_annotation"] is False
    assert trace["vertical_export"]["cross_section_merge_on_source"] is False
    assert sum(
        row["kind"] == "VERTICAL_SECTION"
        for row in trace["rows"]
    ) == 2
    drawing_rows = [row for row in trace["rows"] if row["kind"] == "DRAWING"]
    assert len(drawing_rows) == 2
    assert {
        row["vertical_section_label_handle"]
        for row in drawing_rows
    } == {"T1", "T2"}
    grouped_row = next(
        row for row in drawing_rows
        if row["vertical_section_label_handle"] == "T1"
    )
    assert grouped_row["row"] == 8
    assert grouped_row["source_group_id"] == "verticals!B8"
    assert grouped_row["row_granularity"] == (
        "PRODUCTION_IDENTITY_WITHIN_VERTICAL_SECTION"
    )
    assert grouped_row["source_annotation_count"] == 2
    assert grouped_row["count_formula_terms"] == [25, 25]
    assert grouped_row["count_total"] == 50
    assert grouped_row["annotation_handles"] == ["A1", "A2"]

    annotation_map = {
        row["annotation_handle"]: row
        for row in trace["annotation_position_map"]
    }
    assert annotation_map["A1"]["source_row"] == 8
    assert annotation_map["A2"]["source_row"] == 8
    assert annotation_map["A3"]["source_row"] == 10
    assert annotation_map["A1"]["source_group_id"] == "verticals!B8"
    assert annotation_map["A1"]["source_group_annotation_count"] == 2

    total_row = next(row for row in trace["rows"] if row["kind"] == "TOTAL")
    assert total_row["source_cells"] == ["'verticals'!B8", "'verticals'!B10"]
    assert total_row["source_group_ids"] == ["verticals!B8", "verticals!B10"]
    assert total_row["source_group_count"] == 2
    assert total_row["annotation_count"] == 3
    assert total_row["count_formula_terms"] == [25, 25, 50]
    assert total_row["count_total"] == 100
    assert total_row["annotation_handles"] == ["A1", "A2", "A3"]
